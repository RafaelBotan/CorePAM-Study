from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

PLACEHOLDER_COLOR = RGBColor(0xC0, 0x00, 0x00)

def add_placeholder_para(text_parts):
    p = doc.add_paragraph()
    for text, is_ph in text_parts:
        run = p.add_run(text)
        if is_ph:
            run.bold = True
            run.font.color.rgb = PLACEHOLDER_COLOR
    return p

def add_para(text, bold=False, align=None, size=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    return p

# ---- LETTERHEAD PLACEHOLDER ----
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('[ INSTITUTION LETTERHEAD — Universidade de Brasília (UnB) / School of Medicine / Postgraduate Program — official paper with logo, address, phone, email ]')
run.bold = True
run.italic = True
run.font.color.rgb = PLACEHOLDER_COLOR

doc.add_paragraph()

# ---- DATE ----
add_para('Date: May 11, 2026')

doc.add_paragraph()

# ---- ADDRESSEE ----
add_para('To: APC Discount and Waiver Service')
add_para('Springer Nature')

doc.add_paragraph()

# ---- REFERENCE BLOCK ----
add_para('Re: APC Discount/Waiver Request', bold=True)
add_para('Ticket ID: 48e80b4e-1115-4c71-b38f-cfd034859b43')
add_para('Manuscript: "Corepam: a 24-gene PAM50-derived expression score with cross-platform external validation for breast cancer prognosis"')
add_para('DOI: 10.1186/s13058-026-02298-5')
add_para('Journal: Breast Cancer Research')
add_para('Corresponding author: Dr. Rafael de Negreiros Botan')

doc.add_paragraph()

add_para('To Whom It May Concern,')

doc.add_paragraph()

# ---- OPENING ----
p = doc.add_paragraph()
p.add_run('I am writing in my capacity as Coordinator of the Postgraduate Program of the School of Medicine, Universidade de Brasília (UnB), to formally confirm the funding status of Dr. Rafael de Negreiros Botan in relation to the above-referenced manuscript.')

add_para('I hereby certify that:')

def add_numbered(num, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.75)
    p.add_run(f'{num}. {text}')

add_numbered(1, 'Dr. Rafael de Negreiros Botan is currently enrolled as a doctoral candidate (doutorando) in the Postgraduate Program of the School of Medicine, Universidade de Brasília (UnB).')

add_numbered(2, 'The research presented in the manuscript "Corepam: a 24-gene PAM50-derived expression score with cross-platform external validation for breast cancer prognosis" was conducted without any external research grant, institutional research funding, industry sponsorship, or dedicated publication budget.')

add_numbered(3, 'Neither Dr. Botan nor this institution has any allocated funds — departmental, programmatic, or grant-based — available to cover article processing charges (APCs) for this manuscript.')

add_numbered(4, 'No open-access publication fund, library agreement, or transformative/read-and-publish agreement covering Breast Cancer Research is available to the author through this institution.')

add_numbered(5, 'I am not an author or co-author of the manuscript referenced above, and I have no conflict of interest in providing this statement.')

doc.add_paragraph()

# ---- CLOSING ----
add_para("Given the absence of any source of funding to cover the APC, I respectfully support Dr. Botan's request for an additional discount or full waiver, so that this accepted manuscript on breast cancer — a topic of significant clinical and public-health relevance — can be made openly accessible.")

doc.add_paragraph()
add_para('Should you require any additional information or verification, please do not hesitate to contact me at the details below.')

doc.add_paragraph()
add_para('Sincerely,')

doc.add_paragraph()
doc.add_paragraph()
add_para('_______________________________________')

p = doc.add_paragraph()
r = p.add_run('[FULL NAME OF SIGNATORY]')
r.bold = True
r.font.color.rgb = PLACEHOLDER_COLOR

add_para('Coordinator of the Postgraduate Program')
add_para('School of Medicine, Universidade de Brasília (UnB)')

p = doc.add_paragraph()
p.add_run('Email: ')
r = p.add_run('[INSTITUTIONAL EMAIL]')
r.bold = True
r.font.color.rgb = PLACEHOLDER_COLOR

p = doc.add_paragraph()
p.add_run('Phone: ')
r = p.add_run('[INSTITUTIONAL PHONE]')
r.bold = True
r.font.color.rgb = PLACEHOLDER_COLOR

p = doc.add_paragraph()
r = p.add_run('[OFFICIAL STAMP, if applicable]')
r.bold = True
r.font.color.rgb = PLACEHOLDER_COLOR

# ---- INSTRUCTIONS PAGE ----
doc.add_page_break()

add_para('INSTRUCTIONS FOR USE (remove this page before submission)', bold=True, size=12)
doc.add_paragraph()

add_para('Remaining placeholders (in bold red):', bold=True)
items = [
    '[INSTITUTION LETTERHEAD] — print the letter on the official UnB / School of Medicine / Postgraduate Program letterhead.',
    '[FULL NAME OF SIGNATORY] — full name of the current Coordinator of the Postgraduate Program.',
    '[INSTITUTIONAL EMAIL] — official institutional email of the Coordinator.',
    '[INSTITUTIONAL PHONE] — official institutional phone of the Coordinator/Program.',
    '[OFFICIAL STAMP] — official stamp of the Postgraduate Program, if available.',
]
for it in items:
    doc.add_paragraph(it, style='List Bullet')

doc.add_paragraph()
add_para('Springer Nature requirements (must be respected):', bold=True)
reqs = [
    'Letter must be printed on official institutional letterhead ("headed paper").',
    'Letter must be physically signed (not just typed). Official stamp recommended.',
    'The signatory must NOT be an author or co-author of the manuscript. The senior co-author (professor) cannot sign.',
    'Scan the signed letter (PDF) and upload it directly to the Springer ticket.',
]
for r in reqs:
    doc.add_paragraph(r, style='List Bullet')

out_path = r'Y:/Doutorado Botan/Estudos/CorePAM_Study_accepted/submission/APC_Waiver_Letter_UnB.docx'
doc.save(out_path)
print(f'Saved: {out_path}')
