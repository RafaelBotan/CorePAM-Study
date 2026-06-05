"""Build a clean PDF of the CorePAM manuscript final accepted version for the banca."""
from copy import deepcopy
from pathlib import Path
import subprocess
import shutil

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = Path(__file__).parent
SRC = HERE / "CorePAM_manuscript_v3_review.docx"
OUT_DOCX = HERE / "CorePAM_FINAL_ACCEPTED_BCR.docx"
OUT_PDF = HERE / "CorePAM_FINAL_ACCEPTED_BCR.pdf"
SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"


def _shade(paragraph, color_hex):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    pPr.append(shd)


def _border(paragraph, color_hex):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "12")
        b.set(qn("w:space"), "6")
        b.set(qn("w:color"), color_hex)
        pBdr.append(b)
    pPr.append(pBdr)


def add_banner(doc: Document) -> None:
    title_p = doc.paragraphs[0]
    new_p = OxmlElement("w:p")
    title_p._p.addprevious(new_p)
    p1 = doc.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p1.add_run("FINAL ACCEPTED MANUSCRIPT")
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    _shade(p1, "0B5394")

    new_p2 = OxmlElement("w:p")
    title_p._p.addprevious(new_p2)
    p2 = doc.paragraphs[1]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(
        "Accepted for publication in "
    )
    r2.font.size = Pt(11)
    r2b = p2.add_run("Breast Cancer Research")
    r2b.italic = True
    r2b.bold = True
    r2b.font.size = Pt(11)
    r2c = p2.add_run(" (BMC, Springer Nature)")
    r2c.font.size = Pt(11)
    _shade(p2, "E8F0FA")
    _border(p2, "0B5394")

    new_p3 = OxmlElement("w:p")
    title_p._p.addprevious(new_p3)
    p3 = doc.paragraphs[2]
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run(
        "Acceptance date: 27 April 2026   ·   Submission ID: c2980ad5-9b3d-41a2-bf47-f093c5dc8c5f   ·   Version: v3.0"
    )
    r3.font.size = Pt(9)
    r3.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    _shade(p3, "E8F0FA")
    _border(p3, "0B5394")

    new_p4 = OxmlElement("w:p")
    title_p._p.addprevious(new_p4)
    p4 = doc.paragraphs[3]
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r4 = p4.add_run(
        "This is the author's final accepted manuscript, prior to journal copy-editing and typesetting. "
        "DOI will be assigned during the publishing/rights stage."
    )
    r4.italic = True
    r4.font.size = Pt(9)
    r4.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def main():
    assert SRC.exists(), f"missing source: {SRC}"
    doc = Document(str(SRC))
    add_banner(doc)
    doc.save(str(OUT_DOCX))
    print(f"wrote {OUT_DOCX} ({OUT_DOCX.stat().st_size/1024:.0f} KB)")

    if OUT_PDF.exists():
        OUT_PDF.unlink()
    cmd = [
        SOFFICE,
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        str(HERE),
        str(OUT_DOCX),
    ]
    print("running:", " ".join(cmd))
    res = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
    print("stdout:", res.stdout[-1000:])
    print("stderr:", res.stderr[-1000:])
    print("return:", res.returncode)
    if OUT_PDF.exists():
        print(f"wrote {OUT_PDF} ({OUT_PDF.stat().st_size/1024/1024:.2f} MB)")
    else:
        raise SystemExit("PDF not produced")


if __name__ == "__main__":
    main()
